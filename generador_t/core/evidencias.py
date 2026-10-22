"""Evidencias SOX: pantallazos de las Ts + hoja 'Integridad y Exactitud'.

Refactor de pantallazosTs.py (captura de pantallazos) y evidencias_sox.ipynb
(conversion csv->xlsx + construccion de la hoja de integridad y exactitud
contra la plantilla), parametrizado con Params en vez de valores quemados.
"""
from __future__ import annotations

import time
from copy import copy
from pathlib import Path
from typing import Callable, List, Optional

import openpyxl as opx
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from selenium import webdriver
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

from . import paths
from .params import Params

Logger = Callable[[str], None]
# progreso(actual, total, texto): avance del paso para la barra de la app.
Progreso = Callable[[int, int, str], None]


def _sin_progreso(actual: int, total: int, texto: str) -> None:
    pass

ZOOM_POR_T = {1: 110, 2: 70, 3: 69, 4: 70, 5: 55, 6: 69, 7: 65, 8: 60, 9: 75, 10: 75, 11: 59}
REPETICIONES = {"T1": 1, "T2": 3, "T4": 2, "T5": 1, "T7": 1, "T9": 1, "T10": 1, "T11": 1}
COLUMNAS = [opx.utils.get_column_letter(i) for i in range(1, 79)]

XPATH_BOTON_LOGIN = '//*[@id="lower"]/button'


def _urls_evidencia(periodo: str, ramo: str, compania: str) -> List[str]:
    base = (
        "https://admriesgo.suranet.com/osiris/servlet/suramericana.rentasv.generico.consultas.Consulta"
        "?h_parametros=true&h_type=html&DSRECURSO={recurso}&CDTIPO_RECURSO=A"
        "{rango_periodo}&CERTIFICADO=&{poliza}&RAMO={ramo}&COMPANIA={compania}"
        "&CODIGO_PLAN=&CODIGO_PRODUCTO=&ESTADO=P&NMCOTIZACION_INICIAL=&NMCOTIZACION_FINAL="
        "&TIPOS_ARCHIVO=&SNENCABEZADO="
    )
    recursos_poliza = [
        ("Vector_Tasas_Crecimiento.sql", "POLIZA_INICIAL=&POLIZA_FINAL="),
        ("estado_amparos_cert.sql", "NMPOLIZA_INICIAL=&NMPOLIZA_FINAL="),
        ("estado_beneficios_cert_mod.sql", "NMPOLIZA_INICIAL=&NMPOLIZA_FINAL="),
        ("estado_gastos_cert.sql", "NMPOLIZA_INICIAL=&NMPOLIZA_FINAL="),
        ("estado_modelos.sql", "NMPOLIZA_INICIAL=&NMPOLIZA_FINAL="),
        ("estado_participacion_cert.sql", "NMPOLIZA_INICIAL=&NMPOLIZA_FINAL="),
        ("estado_seguros_cert_mod.sql", "POLIZA_INICIAL=&POLIZA_FINAL="),
        ("estado_beneficiarios_cert_mod.sql", "NMPOLIZA_INICIAL=&NMPOLIZA_FINAL="),
        ("estado_seguros_benef_cert_mod.sql", "POLIZA_INICIAL=&POLIZA_FINAL="),
        ("estado_crecimiento_mod.sql", "POLIZA_INICIAL=&POLIZA_FINAL="),
    ]
    urls = [
        base.format(
            recurso=recurso, rango_periodo=f"&PERIODO={periodo}", poliza=poliza,
            ramo=ramo, compania=compania,
        )
        for recurso, poliza in recursos_poliza
    ]
    urls.append(
        "https://admriesgo.suranet.com/osiris/servlet/suramericana.rentasv.generico.consultas.Consulta"
        f"?h_parametros=true&h_type=html&DSRECURSO=Prod_Poliza_Periodo_v7.sql&CDTIPO_RECURSO=A"
        f"&O_PERIODO_DESDE={periodo}&O_PERIODO_HASTA={periodo}&O_CDRAMO={ramo}&O_CDCOMPANIA={compania}"
        "&TIPOS_ARCHIVO=&SNENCABEZADO="
    )
    return urls


def crear_chrome_driver() -> webdriver.Chrome:
    from webdriver_manager.chrome import ChromeDriverManager

    options = Options()
    options.add_argument("--start-maximized")
    options.add_argument("--disable-notifications")
    options.add_argument("--disable-infobars")
    options.add_experimental_option("excludeSwitches", ["enable-automation"])
    options.add_experimental_option("useAutomationExtension", False)
    return webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)


def _wait_dom_interactive(driver, timeout: int = 60) -> None:
    WebDriverWait(driver, timeout).until(
        lambda d: d.execute_script("return document.readyState") in ("interactive", "complete")
    )


def _click_seguro(driver, locator) -> None:
    element = WebDriverWait(driver, 20).until(EC.element_to_be_clickable(locator))
    try:
        element.click()
    except Exception:
        driver.execute_script("arguments[0].click();", element)


def _send_keys_seguro(driver, locator, keys) -> None:
    element = WebDriverWait(driver, 20).until(EC.visibility_of_element_located(locator))
    try:
        element.clear()
    except Exception:
        pass
    element.send_keys(keys)


def capturar_evidencias_sox(
    params: Params,
    usuario: str,
    contrasena: str,
    driver: Optional[webdriver.Chrome] = None,
    log: Logger = print,
    progreso: Progreso = _sin_progreso,
) -> None:
    """Toma los pantallazos de evidencia (arriba/abajo, +derecha para T8/T11)."""
    periodo = params.fecha_valoracion_input  # para las consultas de Osiris
    carpeta = paths.imagenes_evidencia_dir(params.clave_periodo)
    paths.asegurar_directorio(carpeta)

    urls = _urls_evidencia(periodo, params.ramo, params.compania)

    propio = driver is None
    if propio:
        driver = crear_chrome_driver()

    try:
        for cont, url in enumerate(urls, start=1):
            try:
                progreso(cont - 1, len(urls), f"T{cont} de {len(urls)}: abriendo consulta")
                log(f"[{cont}/{len(urls)}] Accediendo a evidencia T{cont}...")
                driver.get(url)
                _wait_dom_interactive(driver, 60)

                try:
                    WebDriverWait(driver, 3).until(EC.alert_is_present())
                    driver.switch_to.alert.accept()
                except Exception:
                    pass

                page_src = driver.page_source.lower()
                if ("username" in page_src) or ("password" in page_src) or ("idp" in driver.current_url.lower()):
                    if driver.find_elements(By.ID, "username"):
                        _send_keys_seguro(driver, (By.ID, "username"), usuario)
                    if driver.find_elements(By.ID, "password"):
                        _send_keys_seguro(driver, (By.ID, "password"), contrasena)
                    try:
                        _click_seguro(driver, (By.XPATH, XPATH_BOTON_LOGIN))
                    except Exception:
                        for loc in [
                            (By.CSS_SELECTOR, 'button[type="submit"]'),
                            (By.XPATH, "//button[contains(.,'Ingresar') or contains(.,'Acceder') or contains(.,'Entrar')]"),
                            (By.XPATH, "//input[@type='submit']"),
                        ]:
                            if driver.find_elements(*loc):
                                _click_seguro(driver, loc)
                                break
                    time.sleep(1.5)
                    if "idp/mfa" in driver.current_url.lower():
                        log("Esperando validacion MFA (10s)...")
                        time.sleep(10)
                    _wait_dom_interactive(driver, 60)

                try:
                    WebDriverWait(driver, 30).until(
                        EC.visibility_of_element_located((By.XPATH, "//table[.//tr]"))
                    )
                except TimeoutException:
                    WebDriverWait(driver, 30).until(
                        EC.presence_of_element_located((By.XPATH, "//table"))
                    )

                log(f"T{cont}: tomando pantallazos (arriba/abajo{'/derecha' if cont in (8, 11) else ''})...")
                time.sleep(1)
                driver.execute_script(f"document.body.style.zoom='{ZOOM_POR_T[cont]}%'")
                time.sleep(1)
                driver.maximize_window()
                time.sleep(1.5)

                driver.execute_script("window.scrollTo({left:0, top:0, behavior:'instant'})")
                ruta_arriba = carpeta / f"T{cont}_arriba.png"
                time.sleep(0.8)
                driver.save_screenshot(str(ruta_arriba))

                driver.execute_script("window.scrollTo({left:0, top:document.body.scrollHeight, behavior:'instant'})")
                ruta_abajo = carpeta / f"T{cont}_abajo.png"
                time.sleep(0.8)
                driver.save_screenshot(str(ruta_abajo))

                rutas_con_fecha = [ruta_arriba, ruta_abajo]

                if cont in (8, 11):
                    driver.execute_script("window.scrollTo({left:document.body.scrollWidth, top:0, behavior:'instant'})")
                    ruta_arriba_der = carpeta / f"T{cont}_arriba_der.png"
                    time.sleep(0.8)
                    driver.save_screenshot(str(ruta_arriba_der))

                    driver.execute_script("window.scrollTo({left:document.body.scrollWidth, top:document.body.scrollHeight, behavior:'instant'})")
                    ruta_abajo_der = carpeta / f"T{cont}_abajo_der.png"
                    time.sleep(0.8)
                    driver.save_screenshot(str(ruta_abajo_der))

                for ruta in rutas_con_fecha:
                    with Image.open(ruta) as img:
                        ancho, alto = img.size
                        img_editable = Image.new("RGB", (ancho, alto + 50), (255, 255, 255))
                        img_editable.paste(img, (0, 50))
                        draw = ImageDraw.Draw(img_editable)
                        draw.text((10, 10), time.strftime("Fecha: %Y-%m-%d %H:%M:%S"), fill="black", font=ImageFont.load_default())
                        img_editable.save(ruta)

                log(f"Pantallazos de evidencia T{cont} listos.")
                progreso(cont, len(urls), f"T{cont} lista")

            except Exception as e:
                log(f"[ERROR] T{cont}: {e}")
                raise
    finally:
        if propio:
            driver.quit()


def convertir_csv_a_xlsx(fecha: str, log: Logger = print, progreso: Progreso = _sin_progreso) -> None:
    """Convierte T1..T11.csv (descargados) a xlsx via COM de Excel, en la misma carpeta."""
    import pythoncom

    from .excel_com import abrir_excel, cerrar_excel, explicar_error

    carpeta = paths.ts_mod_dir(fecha)
    archivos_csv = sorted(carpeta.glob("*.csv"))

    # Streamlit corre el script en un hilo que no es el principal; COM exige
    # que cada hilo que lo use llame CoInitialize (si no, sale el error
    # "-2147221008, no se ha llamado a CoInitialize").
    pythoncom.CoInitialize()
    try:
        excel, pid = abrir_excel()
        wb = None
        try:
            if not archivos_csv:
                raise FileNotFoundError(f"No hay csv para convertir en {carpeta}. Primero descarga las Ts.")
            for n, csv_path in enumerate(archivos_csv, start=1):
                xlsx_path = csv_path.with_suffix(".xlsx")
                progreso(n - 1, len(archivos_csv), f"{csv_path.name} ({n} de {len(archivos_csv)})")
                log(f"Convirtiendo {csv_path.name} -> {xlsx_path.name}")
                wb = excel.Workbooks.Open(str(csv_path))
                try:
                    wb.SaveAs(str(xlsx_path), FileFormat=51)
                    wb.Close()
                    wb = None
                except Exception as e:
                    raise RuntimeError(f"No se pudo convertir {csv_path.name}: {explicar_error(e)}") from e
                progreso(n, len(archivos_csv), f"{xlsx_path.name} listo")
        finally:
            cerrar_excel(excel, wb, pid, log)
        log("Conversion de csv a xlsx completada.")
    finally:
        pythoncom.CoUninitialize()


def _analizar_estructura(fecha: str, numero_t: int) -> List[List[str]]:
    df = pd.read_excel(paths.ts_mod_dir(fecha) / f"T{numero_t}.xlsx")
    num_filas, num_columnas = df.shape

    if numero_t == 8:
        polizas_unicas = df["POLIZA_COTIZACION"].unique()
        primeros_3 = polizas_unicas[:3]
        ultimos_2 = polizas_unicas[-2:]
        cuenta_primeros = sum((df["POLIZA_COTIZACION"] == v).sum() for v in primeros_3)
        cuenta_ultimos = sum((df["POLIZA_COTIZACION"] == v).sum() for v in ultimos_2)
        registros = list(range(2, 2 + cuenta_primeros)) + list(range(num_filas + 2 - cuenta_ultimos, num_filas + 2))
    else:
        rep = REPETICIONES[f"T{numero_t}"]
        registros = list(range(2, 2 + 3 * rep)) + list(range(num_filas + 2 - 2 * rep, num_filas + 2))

    columnas = COLUMNAS[:num_columnas]
    return [[col + str(fila) for col in columnas] for fila in registros]


def _generar_integridad_t(fecha: str, numero_t: int, log: Logger) -> None:
    plantilla = opx.load_workbook(paths.plantilla_path())
    hoja_plantilla = plantilla[f"T{numero_t}"]

    libro_t = opx.load_workbook(paths.ts_mod_dir(fecha) / f"T{numero_t}.xlsx")
    hoja_sox = libro_t.create_sheet(title="Integridad y Exactitud")

    for row in hoja_plantilla.iter_rows():
        for cell in row:
            nueva = hoja_sox.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                nueva.font = copy(cell.font)
                nueva.border = copy(cell.border)
                nueva.fill = copy(cell.fill)
                nueva.number_format = copy(cell.number_format)
                nueva.protection = copy(cell.protection)
                nueva.alignment = copy(cell.alignment)

    if numero_t in (1, 2, 4, 5, 7, 8, 9, 10, 11):
        estructura = _analizar_estructura(fecha, numero_t)
        fila_inicio, columna_inicio = 4, 19
        if numero_t in (8, 11):
            columna_inicio = 33
        for i, fila in enumerate(estructura):
            for j, formula in enumerate(fila):
                hoja_sox.cell(row=fila_inicio + i, column=columna_inicio + j).value = f"='T{numero_t}'!{formula}"

    carpeta_evidencia = paths.imagenes_evidencia_dir(fecha)
    img_arriba = opx.drawing.image.Image(carpeta_evidencia / f"T{numero_t}_arriba.png")
    img_abajo = opx.drawing.image.Image(carpeta_evidencia / f"T{numero_t}_abajo.png")
    img_arriba.width, img_arriba.height = 1100, 500
    img_abajo.width, img_abajo.height = 1100, 500
    hoja_sox.add_image(img_arriba, "A1")
    if numero_t in (1, 2, 4, 5, 7, 8, 9, 10, 11):
        hoja_sox.add_image(img_abajo, "A28")

    if numero_t in (8, 11):
        img_arriba_der = opx.drawing.image.Image(carpeta_evidencia / f"T{numero_t}_arriba_der.png")
        img_abajo_der = opx.drawing.image.Image(carpeta_evidencia / f"T{numero_t}_abajo_der.png")
        for img in (img_arriba, img_abajo, img_arriba_der, img_abajo_der):
            img.width, img.height = 1000, 500
        hoja_sox.add_image(img_arriba_der, "Q1")
        hoja_sox.add_image(img_abajo_der, "Q28")

    if numero_t == 11:
        img_parametros = opx.drawing.image.Image(paths.imagenes_parametros_dir(fecha) / "t11_sox_parametros.png")
        img_parametros.width, img_parametros.height = 1100, 300
        hoja_sox.add_image(img_parametros, "A55")

    destino = paths.asegurar_directorio(paths.archivos_t_final_dir(fecha))
    libro_t.save(destino / f"T{numero_t}.xlsx")
    log(f"Integridad y Exactitud de T{numero_t} generada.")


def generar_integridad_y_exactitud(fecha: str, log: Logger = print, progreso: Progreso = _sin_progreso) -> None:
    """Para T1..T11: crea la hoja 'Integridad y Exactitud' contra la plantilla
    y guarda el resultado final en archivos_T/{fecha}/T{n}.xlsx."""
    for numero_t in range(1, 12):
        progreso(numero_t - 1, 11, f"T{numero_t} de 11")
        log(f"T{numero_t}: armando hoja Integridad y Exactitud...")
        _generar_integridad_t(fecha, numero_t, log)
        progreso(numero_t, 11, f"T{numero_t} lista")
